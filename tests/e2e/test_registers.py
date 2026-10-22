"""E2E: bulk invoice + voucher registers — lazy view, filter, export, print."""
import io
import json
import os
import re

import openpyxl

BASE_URL = "http://127.0.0.1:" + os.environ.get("E2E_PORT", "5050")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _delayed_register_fetch(page):
    """Hold the bulk-document chunk request long enough to test Cancel.

    The page paints its first twenty invoices itself, then asks for the rest;
    the delay lets the test press Cancel while that request is still open,
    exactly like a slow serverless response on the deployed app.
    """
    page.add_init_script("""
      (function(){
        var orig = window.fetch;
        window.fetch = function(url, opts){
          if (String(url).indexOf('/invoicing/registers/data') !== -1) {
            return new Promise(function(resolve, reject){
              var timer = setTimeout(function(){
                orig.call(window, url, opts).then(resolve, reject);
              }, 3000);
              if (opts && opts.signal) {
                opts.signal.addEventListener('abort', function(){
                  clearTimeout(timer);
                  reject(new DOMException('Aborted', 'AbortError'));
                });
              }
            });
          }
          return orig.apply(this, arguments);
        };
      })();
    """)


def _create_draft_invoices(page, count):
    """Create unapproved invoices through the real save endpoint.

    They exist only for the length of one test: a bulk preview needs more
    than the twenty pages the first response carries, and creating them
    through the route keeps numbering and tenancy the same as production.
    """
    customers = page.request.get(
        f"{BASE_URL}/inventory/invoices/api/customers?q=Solar").json()
    assert customers, "the seeded Solar customer is missing"
    customer_id = customers[0]["id"]
    ids = []
    for _ in range(count):
        payload = {
            "invoice_number": "(auto)", "customer_id": customer_id,
            "invoice_date": "2026-06-15", "due_date": "2026-07-15",
            "discount_mode": "general", "charges_mode": "general",
            "tax_mode": "general", "items": [],
            "subtotal": 0, "total_discount": 0, "total_charges": 0,
            "total_tax": 0, "total_further_tax": 0,
            "total_withholding_tax": 0, "total_amount": 0,
            "action": "save",
        }
        resp = page.request.post(
            f"{BASE_URL}/inventory/invoices/save", data=json.dumps(payload),
            headers={"Content-Type": "application/json"})
        assert resp.ok, resp.text()
        ids.append(resp.json()["id"])
    return ids


def _delete_invoices(page, ids):
    for invoice_id in ids:
        resp = page.request.post(
            f"{BASE_URL}/inventory/invoices/delete/{invoice_id}")
        assert resp.ok, resp.text()


def _create_taxed_draft(page):
    """One real combined-tax sales invoice, so the workbook has lines whose
    tax must add up to the invoice header (the export used to show 0% tax on
    every line of a Combined-mode invoice)."""
    page.goto(f"{BASE_URL}/inventory/invoices/")
    page.wait_for_load_state("networkidle")
    page.locator("#gateContinue").click()
    page.wait_for_timeout(250)
    page.locator("#customerSearch").fill("Solar")
    page.locator("#customerDropdown .ac-it").first.click()
    page.wait_for_timeout(120)
    page.locator('#itemsBody [data-col="code"]').first.fill("CBL-SOL-4MM")
    page.wait_for_function(
        "document.querySelector('#itemsBody tr').dataset.productId !== ''")
    page.locator('#itemsBody [data-col="quantity"]').first.fill("2")
    page.locator('#itemsBody [data-col="unit_price"]').first.fill("180")
    page.locator("#globalSalesTaxPct").fill("18")
    page.locator("#globalSalesTaxPct").dispatch_event("input")
    page.keyboard.press("Control+s")
    page.wait_for_function(
        "() => document.getElementById('invoiceId').value !== ''")


class TestInvoiceRegister:
    def test_view_does_not_autoload(self, admin_page):
        admin_page.goto(f"{BASE_URL}/invoicing/registers/")
        admin_page.wait_for_load_state("networkidle")
        body = admin_page.locator("body").inner_text()
        assert "Invoice Register" in body
        for label in ("Month", "Financial Year", "Custom Range"):
            assert label in body
        assert "Load Invoices" in body
        assert "press View" in body
        # Nothing fetched yet: no rows, no totals.
        assert admin_page.locator("#regBody tr").count() == 0

    def test_view_button_loads_the_table(self, admin_page):
        admin_page.goto(
            f"{BASE_URL}/invoicing/registers/?doc=sales&mode=month&month=2026-01")
        admin_page.wait_for_load_state("networkidle")
        admin_page.click("#regLoadBtn")
        admin_page.wait_for_selector("#regLoaded:not([style*='none'])",
                                     timeout=30000)
        body = admin_page.locator("body").inner_text()
        assert "TOTAL" in body
        assert "January 2026" in body

    def test_purchase_register_loads(self, admin_page):
        admin_page.goto(f"{BASE_URL}/invoicing/registers/?doc=purchase")
        admin_page.wait_for_load_state("networkidle")
        body = admin_page.locator("body").inner_text()
        assert "Purchase Invoice Register" in body
        assert "Supplier" in body

    def test_month_filter_labels_the_view(self, admin_page):
        admin_page.goto(
            f"{BASE_URL}/invoicing/registers/?doc=sales&mode=month&month=2026-01")
        admin_page.wait_for_load_state("networkidle")
        assert "January 2026" in admin_page.locator("body").inner_text()

    def test_custom_filter_labels_the_view(self, admin_page):
        admin_page.goto(
            f"{BASE_URL}/invoicing/registers/?doc=sales&mode=custom"
            "&from=2026-01-05&to=2026-01-20")
        admin_page.wait_for_load_state("networkidle")
        body = admin_page.locator("body").inner_text()
        assert "05 Jan 2026" in body and "20 Jan 2026" in body

    def test_export_buttons_appear_after_load(self, admin_page):
        admin_page.goto(
            f"{BASE_URL}/invoicing/registers/?doc=sales&mode=month&month=2026-01")
        admin_page.wait_for_load_state("networkidle")
        admin_page.click("#regLoadBtn")
        admin_page.wait_for_selector("#regLoaded:not([style*='none'])",
                                     timeout=30000)
        excel = admin_page.locator("a[href*='fmt=excel']").first
        pdf = admin_page.locator("a[href*='fmt=pdf']").first
        assert "month=2026-01" in (excel.get_attribute("href") or "")
        assert "month=2026-01" in (pdf.get_attribute("href") or "")
        assert admin_page.locator("#regActions button:has-text('Print')").is_visible()

    def test_rows_link_to_full_invoice_view(self, admin_page):
        admin_page.goto(
            f"{BASE_URL}/invoicing/registers/?doc=sales&mode=custom&from=2020-01-01&to=2030-01-01&view=1")
        admin_page.wait_for_load_state("networkidle")
        links = admin_page.locator(".reg-table a[href*='/inventory/invoices/']")
        if links.count():
            assert "View" in (links.first.inner_text() or "")

    def test_excel_export_downloads(self, admin_page):
        resp = admin_page.request.get(
            f"{BASE_URL}/invoicing/registers/export?fmt=excel&doc=sales")
        assert resp.status == 200
        assert XLSX in resp.headers.get("content-type", "")
        assert len(resp.body()) > 1000

    def test_pdf_export_downloads(self, admin_page):
        resp = admin_page.request.get(
            f"{BASE_URL}/invoicing/registers/export?fmt=pdf&doc=purchase")
        assert resp.status == 200
        assert "application/pdf" in resp.headers.get("content-type", "")
        assert resp.body()[:5] == b"%PDF-"

    def test_unknown_format_is_404(self, admin_page):
        resp = admin_page.request.get(
            f"{BASE_URL}/invoicing/registers/export?fmt=docx")
        assert resp.status == 404
        # CSV is a supported register export now.
        csv_resp = admin_page.request.get(
            f"{BASE_URL}/invoicing/registers/export?fmt=csv")
        assert csv_resp.status == 200
        assert csv_resp.headers["content-type"].startswith("text/csv")

    def test_bulk_documents_load_full_invoices(self, admin_page):
        admin_page.goto(
            f"{BASE_URL}/invoicing/registers/documents?doc=sales&mode=custom&from=2020-01-01&to=2030-01-01")
        admin_page.wait_for_load_state("networkidle")
        # Chunked loader finishes: overlay closes, actions appear.
        admin_page.wait_for_selector("body[data-docs='done']", timeout=60000)
        body = admin_page.locator("body").inner_text()
        assert "Sales Invoice Register" in body

    def test_bulk_documents_excel_is_the_invoice_workbook(self, admin_page):
        """Excel from the bulk-print page carries the invoices, not the
        register summary the register page already exports."""
        _create_taxed_draft(admin_page)
        admin_page.goto(
            f"{BASE_URL}/invoicing/registers/documents?doc=sales&mode=custom&from=2020-01-01&to=2030-01-01")
        admin_page.wait_for_selector("body[data-docs='done']", timeout=60000)
        href = admin_page.locator(
            "#docActions a:has-text('Excel')").get_attribute("href") or ""
        assert "/invoicing/registers/documents/export" in href
        assert "fmt=excel" in href

        resp = admin_page.request.get(BASE_URL + href)
        assert resp.status == 200
        assert XLSX in resp.headers.get("content-type", "")
        wb = openpyxl.load_workbook(io.BytesIO(resp.body()), read_only=True,
                                    data_only=True)
        try:
            assert wb.sheetnames[:3] == ["Invoices", "Items", "Charges"]
            inv_ws = wb["Invoices"]
            inv_headers = next(inv_ws.iter_rows(min_row=5, max_row=5,
                                                values_only=True))
            assert "Invoice #" in inv_headers and "Customer" in inv_headers
            assert "Total" in inv_headers
            item_ws = wb["Items"]
            item_headers = next(item_ws.iter_rows(min_row=5, max_row=5,
                                                  values_only=True))
            assert "Description" in item_headers and "SKU" in item_headers
            assert inv_ws.max_row > 5 and item_ws.max_row > 5
            first_inv = next(inv_ws.iter_rows(
                min_row=6, max_row=6, values_only=True))[1]
            item_invoices = [r[0] for r in item_ws.iter_rows(
                min_row=6, values_only=True)]
            assert first_inv in item_invoices

            # The lines' tax must add up to each invoice's own tax, in both
            # Combined and Per-line modes (combined used to show 0% per line).
            inv_tax_col = inv_headers.index("Sales Tax")
            inv_no_col = inv_headers.index("Invoice #")
            inv_vch_col = inv_headers.index("Voucher #")
            header_tax = {}
            for row in inv_ws.iter_rows(min_row=6, values_only=True):
                if row[inv_no_col] and row[0] != "TOTAL":
                    header_tax[(row[inv_no_col], row[inv_vch_col])] = \
                        float(row[inv_tax_col] or 0)
            item_tax_col = item_headers.index("Sales Tax")
            line_tax = {}
            for row in item_ws.iter_rows(min_row=6, values_only=True):
                if row[0]:
                    key = (row[0], row[1])
                    line_tax[key] = round(
                        line_tax.get(key, 0.0) + float(row[item_tax_col] or 0),
                        2)
            assert header_tax
            for key, tax in header_tax.items():
                assert abs(line_tax.get(key, 0.0) - tax) <= 0.01, \
                    (key, tax, line_tax.get(key))
        finally:
            wb.close()

    def test_bulk_documents_preload_chunks_and_survives_cancel(
            self, admin_page):
        """First twenty pages ship in the HTML; the rest arrive 100 at a
        time; Cancel leaves the toolbar usable and whatever already loaded
        printable. The old page fetched 20 at a time behind a full-screen
        overlay, and cancelling left Print All / Back unreachable."""
        url = (f"{BASE_URL}/invoicing/registers/documents?doc=sales"
               "&mode=custom&from=2020-01-01&to=2030-01-01")
        ids = _create_draft_invoices(admin_page, 25)
        try:
            # --- first response already carries the first twenty pages ----
            resp = admin_page.request.get(url)
            assert resp.status == 200
            html = resp.text()
            total = int(re.search(r'id="docCount">(\d+)<', html).group(1))
            assert total > 20, "need more than one response of documents"
            assert 'data-docs="loading"' in html
            assert html.count('<div class="doc">') == 20
            # A browser without JavaScript still gets a way to every page.
            assert "Open every invoice" in html
            assert "view=1" in html

            # --- the rest arrives in hundred-page chunks -------------------
            chunk_urls = []

            def _collect(req):
                if "/invoicing/registers/data?full=1" in req.url:
                    chunk_urls.append(req.url)

            admin_page.on("request", _collect)
            _delayed_register_fetch(admin_page)
            admin_page.goto(url)
            admin_page.wait_for_selector("body[data-docs='loading']")
            assert admin_page.locator("#docList .doc").count() == 20
            assert admin_page.locator("#regProgress.open").is_visible()
            assert not admin_page.locator("#printAllBtn").is_enabled()
            # The progress card sits at the bottom; it must not cover the
            # Back / Excel / Print actions while a slow chunk is loading.
            back = admin_page.locator(
                "#docActions a:has-text('Back to Register')")
            assert back.is_visible()
            assert admin_page.evaluate("""() => {
                const a = [...document.querySelectorAll('#docActions a')]
                  .find(x => x.textContent.includes('Back'));
                const r = a.getBoundingClientRect();
                const hit = document.elementFromPoint(
                  r.left + r.width / 2, r.top + r.height / 2);
                return hit === a || a.contains(hit);
            }""")
            admin_page.wait_for_selector("body[data-docs='done']", timeout=60000)
            assert admin_page.locator("#docList .doc").count() == total
            assert admin_page.locator("#printAllBtn").is_enabled()
            assert chunk_urls, "the remaining documents were never requested"
            assert all("limit=100" in u for u in chunk_urls)
            assert all("offset=20" in u for u in chunk_urls)
            assert all("total=" in u for u in chunk_urls)

            # --- cancel mid-load keeps the first twenty and the toolbar ----
            admin_page.goto(url)
            admin_page.wait_for_selector("body[data-docs='loading']")
            admin_page.click("#regCancel")
            admin_page.wait_for_selector("body[data-docs='cancelled']")
            assert ("Loading stopped. 20 of %d invoices loaded." % total
                    in admin_page.locator("#docStatus").inner_text())
            assert admin_page.locator("#docList .doc").count() == 20
            assert not admin_page.locator("#regProgress").is_visible()
            assert admin_page.locator("#printAllBtn").is_enabled()
            admin_page.click("#docActions a:has-text('Back to Register')")
            admin_page.wait_for_url("**/invoicing/registers/**")

            # --- no-JS fallback renders every page in the HTML -------------
            full = admin_page.request.get(url + "&view=1").text()
            assert 'data-docs="done"' in full
            assert int(re.search(r'id="docCount">(\d+)<',
                                 full).group(1)) == total
            assert full.count('<div class="doc">') == total
        finally:
            _delete_invoices(admin_page, ids)

    def test_bulk_data_clamps_a_chunk_to_hundred(self, admin_page):
        base = (f"{BASE_URL}/invoicing/registers/data?full=1&doc=sales"
                "&mode=custom&from=2020-01-01&to=2030-01-01")
        resp = admin_page.request.get(base + "&limit=500&offset=999999")
        assert resp.status == 200
        assert resp.json()["limit"] == 100
        # The browser hands the count back after the first response, so a
        # long preview does not count the same filtered set on every chunk.
        resp = admin_page.request.get(base + "&limit=1&offset=0&total=987")
        assert resp.json()["total"] == 987


class TestVoucherRegister:
    def test_view_does_not_autoload(self, admin_page):
        admin_page.goto(f"{BASE_URL}/accounting/registers/vouchers")
        admin_page.wait_for_load_state("networkidle")
        body = admin_page.locator("body").inner_text()
        assert "Voucher Register" in body
        assert "Load Vouchers" in body
        assert "Load Vouchers" in (
            admin_page.locator("#regDocsBtn").inner_text() or "")
        assert admin_page.locator("#regBody tr").count() == 0

    def test_view_button_loads_the_table(self, admin_page):
        admin_page.goto(
            f"{BASE_URL}/accounting/registers/vouchers?mode=month&month=2026-01")
        admin_page.wait_for_load_state("networkidle")
        admin_page.click("#regLoadBtn")
        admin_page.wait_for_selector("#regLoaded:not([style*='none'])",
                                     timeout=30000)
        body = admin_page.locator("body").inner_text()
        assert "TOTAL" in body
        assert "Description" in body
        assert "Status" not in admin_page.locator(
            ".voucher-table thead").inner_text()

    def test_type_filter_labels_the_view(self, admin_page):
        admin_page.goto(
            f"{BASE_URL}/accounting/registers/vouchers?vtype=JV")
        admin_page.wait_for_load_state("networkidle")
        assert "Journal Voucher" in admin_page.locator("body").inner_text()

    def test_rows_link_to_full_voucher_view(self, admin_page):
        admin_page.goto(
            f"{BASE_URL}/accounting/registers/vouchers?mode=custom&from=2020-01-01&to=2030-01-01&view=1")
        admin_page.wait_for_load_state("networkidle")
        links = admin_page.locator("a.vnum")
        if links.count():
            href = links.first.get_attribute("href") or ""
            assert "/accounting/vouchers/" in href and "/preview" in href

    def test_excel_export_downloads(self, admin_page):
        resp = admin_page.request.get(
            f"{BASE_URL}/accounting/registers/vouchers/export?fmt=excel")
        assert resp.status == 200
        assert XLSX in resp.headers.get("content-type", "")
        assert len(resp.body()) > 1000

    def test_pdf_export_downloads(self, admin_page):
        resp = admin_page.request.get(
            f"{BASE_URL}/accounting/registers/vouchers/export?fmt=pdf&vtype=JV")
        assert resp.status == 200
        assert "application/pdf" in resp.headers.get("content-type", "")
        assert resp.body()[:5] == b"%PDF-"

    def test_unknown_format_is_404(self, admin_page):
        resp = admin_page.request.get(
            f"{BASE_URL}/accounting/registers/vouchers/export?fmt=docx")
        assert resp.status == 404
        # CSV is a supported register export now.
        csv_resp = admin_page.request.get(
            f"{BASE_URL}/accounting/registers/vouchers/export?fmt=csv")
        assert csv_resp.status == 200
        assert csv_resp.headers["content-type"].startswith("text/csv")

    def test_bulk_documents_load_full_vouchers(self, admin_page):
        chunk_urls = []

        def _collect(req):
            if "/accounting/registers/vouchers/data?full=1" in req.url:
                chunk_urls.append(req.url)

        admin_page.on("request", _collect)
        admin_page.goto(
            f"{BASE_URL}/accounting/registers/vouchers/documents?mode=custom&from=2020-01-01&to=2030-01-01")
        admin_page.wait_for_load_state("networkidle")
        admin_page.wait_for_selector("body[data-docs='done']", timeout=60000)
        body = admin_page.locator("body").inner_text()
        assert "Bulk Vouchers" in body
        # The first request is small so something paints quickly; every
        # later one asks for a hundred vouchers, not twenty.
        if chunk_urls:
            assert "limit=20" in chunk_urls[0]
            assert all("limit=100" in u for u in chunk_urls[1:])
            total = int(admin_page.locator("#docCount").inner_text())
            assert admin_page.locator("#docList .doc").count() == total

    def test_bulk_voucher_data_clamps_and_reuses_total(self, admin_page):
        base = (f"{BASE_URL}/accounting/registers/vouchers/data?full=1"
                "&mode=custom&from=2020-01-01&to=2030-01-01")
        resp = admin_page.request.get(base + "&limit=500&offset=999999")
        assert resp.status == 200
        assert resp.json()["limit"] == 100
        resp = admin_page.request.get(base + "&limit=1&offset=0&total=654")
        assert resp.json()["total"] == 654
