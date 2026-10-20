"""Customers, suppliers, products and units: the forms and the Excel imports.

The imports kept the parsed sheet in the session cookie on the way to the
review screen. A browser drops a cookie over 4 KB, which about forty
realistic customers exceed, so the review screen opened empty and the import
silently did nothing. The review screen is now the upload's own response.

The forms checked nothing on the server (a blank name was saved, a duplicate
SKU reached the unique key as a server error), the batch save truncated
opening stock to whole numbers ("2.5" became 0) and crashed on a number where
it expected text, and a unit could be deleted while products used it.
"""
import json
import os
import random
import string
import tempfile
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

import pytest

_TMP_DB = os.path.join(tempfile.gettempdir(), "erp_master_data_imports.db")
if os.path.exists(_TMP_DB):
    os.remove(_TMP_DB)
os.environ["DATABASE_URL"] = (os.environ.get("TEST_DATABASE_URL")
                              or "sqlite:///" + _TMP_DB.replace("\\", "/"))

from app import app as flask_app  # noqa: E402
from shared.extensions import db  # noqa: E402

CUSTOMER_HEAD = ["Name", "Contact Person", "Email", "Phone", "Mobile", "City", "Tax ID",
                 "Payment Terms", "Credit Limit", "Website", "Address", "Notes"]
SUPPLIER_HEAD = ["Name", "Contact Person", "Email", "Phone", "Mobile", "City", "Tax ID",
                 "Payment Terms", "Website", "Address", "Notes"]
PRODUCT_HEAD = ["SKU", "Name", "Description", "Category", "Unit Price", "Cost Price",
                "Reorder Level", "Opening Stock", "Unit", "HS Code"]


@pytest.fixture(scope="module")
def world():
    c = flask_app.test_client()
    c.get("/")
    from hr_app.models.user import User
    from shared.models.company import Company, CompanyMembership
    from shared.models.base import Role
    from shared.tenancy import set_current_company, unscoped
    from shared.company_setup import provision_company
    from shared.models.company_settings import AccountingPeriod
    from inventory_app.models.customer import InvCustomer
    from inventory_app.models.category import InvCategory
    from inventory_app.models.unit import InvUnit

    with flask_app.app_context():
        user = (User.query.filter_by(is_super_admin=True).order_by(User.id).first()
                or User.query.order_by(User.id).first())
        uid = user.id
        with unscoped():
            comp = Company(name="Imports Co", slug="imports-co-e2e", is_active=True,
                           created_by=uid)
            db.session.add(comp)
            db.session.flush()
            db.session.add(CompanyMembership(
                company_id=comp.id, user_id=uid,
                role_id=Role.query.filter_by(name=Role.ADMIN).first().id,
                status=CompanyMembership.ACTIVE))
            db.session.commit()
            cid = comp.id
        set_current_company(cid)
        provision_company(cid)
        set_current_company(cid)
        for name, abbr in (("Piece", "pcs"), ("Kilogram", "kg"), ("Meter", "m")):
            if not InvUnit.query.filter_by(abbreviation=abbr).first():
                db.session.add(InvUnit(name=name, abbreviation=abbr))
        db.session.add(InvCategory(name="Solar Panels"))
        db.session.add(InvCustomer(name="Existing Customer"))
        # A closed year long before anything else, for the period-lock tests.
        db.session.add(AccountingPeriod(fiscal_year="FY2020", period_name="FY 2020",
                                        start_date=date(2020, 1, 1),
                                        end_date=date(2020, 12, 31),
                                        is_open=False, is_closed=True,
                                        closed_at=datetime(2021, 1, 15)))
        db.session.commit()
    with c.session_transaction() as sess:
        sess["_user_id"] = str(uid)
        sess["_fresh"] = True
        sess["company_id"] = cid
    return {"client": c, "company": cid, "user": uid}


class Ctx:
    def __init__(self, w):
        self.w = w

    def __enter__(self):
        from shared.tenancy import set_current_company
        self.cm = flask_app.app_context()
        self.cm.__enter__()
        set_current_company(self.w["company"])
        return self

    def __exit__(self, *a):
        db.session.remove()
        self.cm.__exit__(*a)


def _rs(k):
    return "".join(random.choice(string.ascii_letters) for _ in range(k))


def _xlsx(head, rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(head)
    for r in rows:
        ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _upload(c, kind, buf):
    return c.post(f"/inventory/{kind}/upload-excel", data={"file": (buf, "sheet.xlsx")},
                  content_type="multipart/form-data")


def _session_cookie_size(resp):
    sizes = [len(v) for k, v in resp.headers.items()
             if k.lower() == "set-cookie" and v.startswith("session=")]
    return max(sizes, default=0)


def _count(model, **kw):
    with Ctx(WORLD):
        return model.query.filter_by(**kw).count()


WORLD = {}


@pytest.fixture(autouse=True)
def _remember(world):
    WORLD.update(world)


# ── Customers / suppliers ────────────────────────────────────────────────────

def _customer_rows(n, tag):
    random.seed(n)
    return [[f"{_rs(8)} Traders {tag}{i}", _rs(10), f"{_rs(6).lower()}@{_rs(5).lower()}.com",
             f"042-{random.randint(1000000, 9999999)}", f"0300-{random.randint(1000000, 9999999)}",
             random.choice(["Lahore", "Karachi", "Multan"]),
             f"{random.randint(1000000, 9999999)}-{random.randint(0, 9)}", "Net 30",
             random.randint(1, 9) * 10000, "", f"Plot {random.randint(1, 999)}, {_rs(7)} Road",
             _rs(12)] for i in range(n)]


def test_sixty_customers_reach_the_review_screen_without_the_cookie(world):
    c = world["client"]
    rows = _customer_rows(60, "A")
    r = _upload(c, "customers", _xlsx(CUSTOMER_HEAD, rows))
    body = r.get_data(as_text=True)
    assert r.status_code == 200, body[:300]
    assert all(row[0] in body for row in rows), "every row is on the review screen"
    # The rows used to be ~6 KB of cookie; now the cookie does not carry them.
    assert _session_cookie_size(r) < 1024, _session_cookie_size(r)

    data = [{"name": row[0], "contact_person": row[1], "email": row[2], "phone": row[3],
             "mobile": row[4], "city": row[5], "tax_id": row[6], "payment_terms": row[7],
             "credit_limit": str(row[8]), "website": row[9], "address": row[10],
             "notes": row[11]} for row in rows]
    r = c.post("/inventory/customers/batch", data={"data": json.dumps(data)})
    assert r.status_code == 302, r.get_data(as_text=True)[:300]
    from inventory_app.models.customer import InvCustomer
    from shared.models.ledger import ChartOfAccount
    with Ctx(world):
        made = InvCustomer.query.filter(InvCustomer.name.like("% Traders A%")).all()
        assert len(made) == 60
        assert ChartOfAccount.query.filter(
            ChartOfAccount.name.in_([m.name for m in made])).count() == 60


def test_upload_flags_problem_rows_instead_of_dropping_them(world):
    c = world["client"]
    rows = [["Good Buyer", "", "good@buyer.com", "", "", "Lahore", "", "Net 90", 1000, "", "", ""],
            ["  existing   customer ", "", "", "", "", "", "", "", 0, "", "", ""],
            ["Good Buyer", "", "", "", "", "", "", "", 0, "", "", ""],
            ["Minus Credit", "", "", "", "", "", "", "", -500, "", "", ""],
            ["Bad Mail", "", "not-an-email", "", "", "", "", "", 0, "", "", ""],
            ["Long Phone", "", "", "0300-1234567, 042-7654321", "", "", "", "", 0, "", "", ""]]
    r = _upload(c, "customers", _xlsx(CUSTOMER_HEAD, rows))
    body = r.get_data(as_text=True)
    assert r.status_code == 200
    assert "5 row(s) need attention" in body
    assert "&#39;existing customer&#39; already exists" in body or "'existing customer' already exists" in body
    assert "appears more than once in this sheet" in body
    assert "Credit limit cannot be negative" in body
    assert "does not look like an email address" in body
    assert "Phone is longer than 20 characters" in body
    assert 'value="Net 90" selected' in body, "an unlisted payment term is kept, not lost"
    assert "Row 3:" in body, "problems name their spreadsheet row"


def test_batch_save_keeps_the_rows_it_could_not_save(world):
    c = world["client"]
    data = [{"name": "Batch Keep One"}, {"name": "Batch Keep Two", "credit_limit": "2,500"},
            {"name": "EXISTING customer"}, {"name": "Batch Keep One"},
            {"name": "Neg", "credit_limit": "-1"}]
    r = c.post("/inventory/customers/batch", data={"data": json.dumps(data)})
    body = r.get_data(as_text=True)
    assert r.status_code == 200
    assert "2 customer(s) saved" in body
    assert "3 row(s) need attention" in body
    assert "appears more than once in these rows" in body
    from inventory_app.models.customer import InvCustomer
    assert _count(InvCustomer, name="Batch Keep Two") == 1
    with Ctx(world):
        assert float(InvCustomer.query.filter_by(name="Batch Keep Two").one().credit_limit) == 2500


def test_batch_save_survives_numbers_and_junk(world):
    c = world["client"]
    r = c.post("/inventory/customers/batch", data={"data": json.dumps(
        [{"name": 123456, "phone": 3001234567, "credit_limit": 100, "notes": None},
         "not an object", {"_delete": True, "name": "Deleted Row"}, {}])})
    assert r.status_code == 302, r.get_data(as_text=True)[:400]
    from inventory_app.models.customer import InvCustomer
    with Ctx(world):
        made = InvCustomer.query.filter_by(name="123456").one()
        assert made.phone == "3001234567"
        assert InvCustomer.query.filter_by(name="Deleted Row").count() == 0
    for junk in ("{not json", json.dumps({"name": "x"}), ""):
        r = c.post("/inventory/customers/batch", data={"data": junk})
        assert r.status_code in (200, 302), junk
    assert c.post("/inventory/customers/batch", data={}).status_code in (200, 302)


def test_customer_form_refuses_blank_duplicate_and_negative(world):
    c = world["client"]
    from inventory_app.models.customer import InvCustomer
    before = _count(InvCustomer)
    r = c.post("/inventory/customers/create", data={"name": "   "})
    assert r.status_code == 200 and "Name is required" in r.get_data(as_text=True)
    r = c.post("/inventory/customers/create", data={})          # no fields at all
    assert r.status_code == 200
    r = c.post("/inventory/customers/create", data={"name": "existing  CUSTOMER",
                                                    "city": "Typed City"})
    body = r.get_data(as_text=True)
    assert "already exists" in body and "Typed City" in body, "what was typed comes back"
    r = c.post("/inventory/customers/create", data={"name": "Neg Limit", "credit_limit": "-5"})
    assert "Credit limit cannot be negative" in r.get_data(as_text=True)
    assert _count(InvCustomer) == before

    r = c.post("/inventory/customers/create", data={"name": "Form Customer", "credit_limit": "7000"})
    assert r.status_code == 302
    with Ctx(world):
        cid = InvCustomer.query.filter_by(name="Form Customer").one().id
    r = c.post(f"/inventory/customers/edit/{cid}", data={"name": "Existing Customer",
                                                         "is_active": "on"})
    assert "already exists" in r.get_data(as_text=True)
    r = c.post(f"/inventory/customers/edit/{cid}", data={"name": "Form Customer Renamed",
                                                         "is_active": "on"})
    assert r.status_code == 302
    with Ctx(world):
        assert InvCustomer.query.get(cid).name == "Form Customer Renamed"
    # The live duplicate hint uses the save's rule.
    assert c.get("/inventory/customers/check?value=existing customer").get_json()["exists"]


def test_suppliers_import_and_form(world):
    c = world["client"]
    rows = [[r[0].replace("Traders", "Supplies"), *r[1:8], *r[9:]] for r in _customer_rows(45, "S")]
    r = _upload(c, "suppliers", _xlsx(SUPPLIER_HEAD, rows))
    body = r.get_data(as_text=True)
    assert r.status_code == 200 and all(row[0] in body for row in rows)
    assert _session_cookie_size(r) < 1024
    data = [dict(zip(("name", "contact_person", "email", "phone", "mobile", "city", "tax_id",
                      "payment_terms", "website", "address", "notes"), row)) for row in rows]
    r = c.post("/inventory/suppliers/batch", data={"data": json.dumps(data)})
    assert r.status_code == 302
    from inventory_app.models.supplier import InvSupplier
    with Ctx(world):
        assert InvSupplier.query.filter(InvSupplier.name.like("% Supplies S%")).count() == 45
    r = c.post("/inventory/suppliers/create", data={"name": rows[0][0].upper()})
    assert r.status_code == 200 and "already exists" in r.get_data(as_text=True)


# ── Products ─────────────────────────────────────────────────────────────────

def test_product_template_has_every_column_the_import_reads(world):
    import openpyxl
    r = world["client"].get("/inventory/products/template")
    ws = openpyxl.load_workbook(BytesIO(r.data)).active
    assert [c.value for c in ws[1]] == PRODUCT_HEAD


def test_fractional_opening_stock_is_kept_with_its_journal(world):
    c = world["client"]
    rows = [["FR-1", "Copper Cable", "4mm", "Solar Panels", 900, 600, 5, 2.5, "kg", "8544.4900"],
            ["FR-2", "Mounting Rail", "", "", 1500, 1000, 0, 12, "Meter", ""]]
    r = _upload(c, "products", _xlsx(PRODUCT_HEAD, rows))
    body = r.get_data(as_text=True)
    assert r.status_code == 200 and "Copper Cable" in body and 'value="2.5"' in body
    assert "need attention" not in body
    data = [{"sku": "fr-1", "name": "Copper Cable", "description": "4mm",
             "category_name": "solar panels", "unit_price": "900", "cost_price": "600",
             "reorder_level": "5", "current_stock": "2.5", "unit": "kg", "hs_code": "8544.4900"},
            {"sku": "FR-2", "name": "Mounting Rail", "unit_price": "1500", "cost_price": "1,000",
             "reorder_level": "0", "current_stock": "12", "unit": "m"}]
    r = c.post("/inventory/products/batch",
               data={"data": json.dumps(data), "opening_date": "2025-07-01"})
    assert r.status_code == 302, r.get_data(as_text=True)[:500]
    from inventory_app.models.product import InvProduct
    from shared import costing, integrity
    from shared.models.ledger import JournalEntry
    with Ctx(world):
        p = InvProduct.query.filter_by(sku="FR-1").one()
        assert p.category_id is not None and p.unit == "kg" and p.hs_code == "8544.4900"
        assert costing.on_hand(p.id) == Decimal("2.5")
        assert costing.stock_value(p.id) == Decimal("1500")
        assert p.current_stock == Decimal("2.5")
        je = JournalEntry.query.filter_by(voucher_type="OPENING", voucher_id=p.id).one()
        assert je.entry_date.date() == date(2025, 7, 1)
        assert not integrity.failures(), integrity.failures()


def test_product_rows_with_bad_values_are_flagged_and_kept(world):
    c = world["client"]
    rows = [["BAD-1", "Free Stock", "", "", 100, 0, 0, 5, "pcs", ""],
            ["BAD-2", "Negative Price", "", "", -10, 5, 0, 0, "pcs", ""],
            ["BAD-3", "Odd Unit", "", "", 10, 5, 0, 0, "crates", ""],
            ["BAD-4", "Odd Category", "", "Nope", 10, 5, 0, 0, "pcs", ""],
            ["BAD-5", "Half Reorder", "", "", 10, 5, 2.5, 0, "pcs", ""],
            ["fr-1", "Copper Cable Again", "", "", 10, 5, 0, 0, "pcs", ""],
            ["BAD-7", "Bad Stock", "", "", 10, 5, 0, "lots", "pcs", ""]]
    r = _upload(c, "products", _xlsx(PRODUCT_HEAD, rows))
    body = r.get_data(as_text=True)
    assert r.status_code == 200
    assert "7 row(s) need attention" in body
    for msg in ("Opening stock needs a cost price", "Sale price cannot be negative",
                "Unit &#39;crates&#39; is not set up", "Category &#39;Nope&#39; not found",
                "Reorder level must be a whole number", "is already used by another product",
                "Opening stock must be a number"):
        assert msg in body, msg
    assert 'value="crates" selected' in body, "an unknown unit is kept, not swapped for the first one"
    data = [{"sku": "BAD-1", "name": "Free Stock", "unit_price": "100", "cost_price": "0",
             "current_stock": "5", "unit": "pcs"}]
    r = c.post("/inventory/products/batch", data={"data": json.dumps(data)})
    assert r.status_code == 200 and "Opening stock needs a cost price" in r.get_data(as_text=True)
    from inventory_app.models.product import InvProduct
    assert _count(InvProduct, sku="BAD-1") == 0


def test_opening_stock_in_a_closed_period_is_refused_cleanly(world):
    c = world["client"]
    from inventory_app.models.product import InvProduct
    data = [{"sku": "CL-1", "name": "Closed Period Item", "unit_price": "10",
             "cost_price": "5", "current_stock": "3", "unit": "pcs"}]
    r = c.post("/inventory/products/batch",
               data={"data": json.dumps(data), "opening_date": "2020-06-30"})
    body = r.get_data(as_text=True)
    assert r.status_code == 200 and "FY 2020" in body and "Closed Period Item" in body
    assert _count(InvProduct, sku="CL-1") == 0
    r = c.post("/inventory/products/create", data={
        "sku": "CL-2", "name": "Closed Form Item", "unit_price": "10", "cost_price": "5",
        "current_stock": "3", "opening_date": "2020-06-30", "unit": "pcs"})
    body = r.get_data(as_text=True)
    assert r.status_code == 200 and "FY 2020" in body
    assert _count(InvProduct, sku="CL-2") == 0
    r = c.post("/inventory/products/create", data={
        "sku": "CL-3", "name": "Bad Date Item", "cost_price": "5", "current_stock": "3",
        "opening_date": "31/06/2025", "unit": "pcs"})
    assert r.status_code == 200 and "not a valid date" in r.get_data(as_text=True)


def test_product_form_refuses_duplicates_without_a_server_error(world):
    c = world["client"]
    from inventory_app.models.product import InvProduct
    for sku, name in (("FR-1", "Different Name"), ("fr-1 ", "Different Name"),
                      ("NEW-9", " copper   cable ")):
        r = c.post("/inventory/products/create", data={"sku": sku, "name": name,
                                                       "unit_price": "1", "unit": "pcs"})
        body = r.get_data(as_text=True)
        assert r.status_code == 200, body[:300]
        assert "already" in body
    assert _count(InvProduct, sku="NEW-9") == 0
    r = c.post("/inventory/products/create", data={})
    assert r.status_code == 200 and "SKU is required" in r.get_data(as_text=True)
    r = c.post("/inventory/products/create", data={"sku": "NEG-1", "name": "Neg",
                                                   "unit_price": "-1", "unit": "pcs"})
    assert "Sale price cannot be negative" in r.get_data(as_text=True)
    r = c.post("/inventory/products/create", data={"sku": "BIG-1", "name": "Big",
                                                   "unit_price": "1e400", "unit": "pcs"})
    assert "too large" in r.get_data(as_text=True)
    with Ctx(world):
        pid = InvProduct.query.filter_by(sku="FR-2").one().id
    r = c.post(f"/inventory/products/edit/{pid}", data={"sku": "FR-1", "name": "Mounting Rail",
                                                        "unit": "m", "is_active": "on"})
    assert r.status_code == 200 and "already used" in r.get_data(as_text=True)


def test_editing_a_product_keeps_a_unit_that_was_switched_off(world):
    c = world["client"]
    from inventory_app.models.product import InvProduct
    from inventory_app.models.unit import InvUnit
    with Ctx(world):
        db.session.add(InvUnit(name="Carton", abbreviation="ctn", is_active=False))
        p = InvProduct(sku="CTN-1", name="Boxed Item", unit="ctn")
        db.session.add(p)
        db.session.commit()
        pid = p.id
    page = c.get(f"/inventory/products/edit/{pid}").get_data(as_text=True)
    assert 'value="ctn" selected' in page
    r = c.post(f"/inventory/products/edit/{pid}", data={"sku": "CTN-1", "name": "Boxed Item",
                                                        "unit": "ctn", "is_active": "on"})
    assert r.status_code == 302
    with Ctx(world):
        assert InvProduct.query.get(pid).unit == "ctn"


def test_quick_adjust_refuses_junk_quantities(world):
    c = world["client"]
    from inventory_app.models.product import InvProduct
    with Ctx(world):
        pid = InvProduct.query.filter_by(sku="FR-2").one().id
    for junk in ("abc", "inf", "1e400"):
        r = c.post(f"/inventory/products/adjust-stock/{pid}", data={"quantity": junk})
        assert r.status_code == 200, junk


# ── Units ────────────────────────────────────────────────────────────────────

def test_units_forms_imports_and_guards(world):
    c = world["client"]
    from inventory_app.models.unit import InvUnit
    from inventory_app.models.product import InvProduct
    r = c.post("/inventory/units/create", data={"name": "", "abbreviation": ""})
    assert r.status_code == 200 and "Name is required" in r.get_data(as_text=True)
    assert c.post("/inventory/units/create", data={}).status_code == 200
    r = c.post("/inventory/units/create", data={"name": "KILOGRAM", "abbreviation": "KG"})
    body = r.get_data(as_text=True)
    assert "already exists" in body
    assert c.post("/inventory/units/quick-create",
                  data={"name": "Kilos", "abbreviation": "Kg"}).get_json()["ok"] is False
    assert c.post("/inventory/units/quick-create",
                  data={"name": "Litre", "abbreviation": "ltr"}).get_json()["ok"] is True

    rows = [["Gram", "g", "Weight"], ["Grams", "G", ""], ["Meter", "mtr", ""], ["", "x", ""]]
    r = _upload(c, "units", _xlsx(["Name", "Abbreviation", "Explanation"], rows))
    body = r.get_data(as_text=True)
    assert r.status_code == 200 and "3 row(s) need attention" in body

    # In use: cannot be deleted; renamed abbreviation moves its products.
    with Ctx(world):
        u = InvUnit.query.filter_by(abbreviation="ltr").one()
        db.session.add(InvProduct(sku="OIL-1", name="Oil", unit="ltr"))
        db.session.commit()
        uid = u.id
    r = c.post(f"/inventory/units/delete/{uid}")
    assert r.status_code == 302
    assert _count(InvUnit, id=uid) == 1
    r = c.post(f"/inventory/units/edit/{uid}", data={"name": "Litre", "abbreviation": "L",
                                                     "is_active": "on"})
    assert r.status_code == 302
    with Ctx(world):
        assert InvProduct.query.filter_by(sku="OIL-1").one().unit == "L"


def test_a_long_failing_batch_does_not_bloat_the_cookie(world):
    """Batch problems are shown on the page, not in a flash message (which
    lives in the cookie and broke it the same way the rows did)."""
    c = world["client"]
    data = [{"name": "Existing Customer"} for _ in range(150)]
    r = c.post("/inventory/customers/batch", data={"data": json.dumps(data)})
    assert r.status_code == 200
    assert _session_cookie_size(r) < 1024
